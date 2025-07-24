from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from datetime import datetime
import io
import logging

logger = logging.getLogger(__name__)

def format_date(date_obj):
    """Formate une date en string"""
    if date_obj:
        return date_obj.strftime('%d/%m/%Y')
    return ''

def format_time(time_obj):
    """Formate une heure en string"""
    if time_obj:
        return time_obj.strftime('%H:%M')
    return ''

def format_duration(duration):
    """Formate une durée en heures et minutes"""
    if duration:
        hours = duration.seconds // 3600 + duration.days * 24
        minutes = (duration.seconds // 60) % 60
        return f"{hours}h{minutes:02d}min"
    return "Non renseignée"

def export_to_excel(employees, pointages=None):
    """Exporte les données en Excel avec heures travaillées"""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Employés"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        
        # En-têtes avec heures travaillées (sans progression)
        headers = [
            'Matricule', 'Nom', 'Prénom', 'Email', 'Département', 'Statut', 
            'Durée travaillée aujourd\'hui', 'Total semaine', 'Total mois',
            'Dernière Connexion'
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
        
        # Données des employés avec heures
        for row, emp in enumerate(employees, 2):
            ws.cell(row=row, column=1, value=emp.matricule)
            ws.cell(row=row, column=2, value=emp.nom)
            ws.cell(row=row, column=3, value=emp.prenom)
            ws.cell(row=row, column=4, value=emp.email)
            ws.cell(row=row, column=5, value=emp.departement)
            ws.cell(row=row, column=6, value="Actif" if emp.is_active else "Inactif")
            
            # Nouvelles colonnes d'heures
            ws.cell(row=row, column=7, value=format_duration(emp.get_duree_travail_jour()))
            ws.cell(row=row, column=8, value=format_duration(emp.get_total_heures_semaine()))
            ws.cell(row=row, column=9, value=format_duration(emp.get_total_heures_mois()))
            ws.cell(row=row, column=10, value=format_date(emp.last_login))
        
        # Ajustement des colonnes
        for col in ws.columns:
            max_length = 0
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 25)
        
        # Si des pointages sont fournis, créer une deuxième feuille
        if pointages:
            ws_pointages = wb.create_sheet(title="Pointages")
            
            # En-têtes des pointages
            headers = ['Date', 'Employé', 'Type', 'Heure', 'Retard']
            for col, header in enumerate(headers, 1):
                cell = ws_pointages.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
            
            # Données des pointages
            for row, point in enumerate(pointages, 2):
                ws_pointages.cell(row=row, column=1, value=format_date(point.date))
                ws_pointages.cell(row=row, column=2, value=f"{point.employe.nom} {point.employe.prenom}")
                ws_pointages.cell(row=row, column=3, value=point.type.capitalize())
                ws_pointages.cell(row=row, column=4, value=format_time(point.heure))
                ws_pointages.cell(row=row, column=5, value="Oui" if point.retard else "Non")
            
            # Ajustement des colonnes
            for col in ws_pointages.columns:
                max_length = 0
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                ws_pointages.column_dimensions[col[0].column_letter].width = min(max_length + 2, 25)
        
        # Sauvegarde dans un buffer
        excel_file = io.BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        return excel_file
    except Exception as e:
        logger.error(f"Erreur lors de l'export Excel: {str(e)}")
        raise

def export_to_pdf(employees, pointages=None):
    """Exporte les données en PDF avec heures travaillées"""
    try:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        
        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1  # Center alignment
        )
        
        # En-tête
        elements.append(Paragraph("Rapport PresenceHub - Heures travaillées", title_style))
        elements.append(Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", 
                                ParagraphStyle('Date', parent=styles['Normal'], alignment=1)))
        elements.append(Spacer(1, 20))
        
        # Tableau des employés avec heures (sans progression)
        data = [['Matricule', 'Nom', 'Prénom', 'Département', 'Durée jour', 'Total sem', 'Total mois', 'Statut']]
        for emp in employees:
            data.append([
                emp.matricule,
                emp.nom,
                emp.prenom,
                emp.departement,
                format_duration(emp.get_duree_travail_jour()),
                format_duration(emp.get_total_heures_semaine()),
                format_duration(emp.get_total_heures_mois()),
                "Actif" if emp.is_active else "Inactif"
            ])
        
        # Style du tableau
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F81BD')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(table)
        
        # Si des pointages sont fournis, ajouter une section pointages
        if pointages:
            elements.append(Spacer(1, 30))
            elements.append(Paragraph("Pointages du jour", title_style))
            elements.append(Spacer(1, 20))
            
            data = [['Date', 'Employé', 'Type', 'Heure', 'Retard']]
            for point in pointages:
                data.append([
                    format_date(point.date),
                    f"{point.employe.nom} {point.employe.prenom}",
                    point.type.capitalize(),
                    format_time(point.heure),
                    "Oui" if point.retard else "Non"
                ])
            
            table = Table(data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F81BD')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]))
            elements.append(table)
        
        # Génération du PDF
        doc.build(elements)
        buffer.seek(0)
        return buffer
    except Exception as e:
        logger.error(f"Erreur lors de l'export PDF: {str(e)}")
        raise